"""appellate_toc_bind.py -- Module B / Unit 11: lock the TRAP 38.1 section
structure into the brief DOCX as content controls, and reattach the TOC when an
edited copy is brought back in.

Two jobs:

1. Build-time wrapping (wrap_section): each section becomes a heading content
   control locked against edit+delete (sdtContentLocked) followed by a body
   content control locked against delete only (sdtLocked, contents editable),
   tagged with the section_key. In OnlyOffice the user can type freely inside a
   body but cannot delete, rename, or reorder the section scaffold. Generated
   sections (TOC / Index of Authorities / Certifications) lock their bodies too.

2. Reattach (reattach_brief_toc): when an externally-edited DOCX is dragged back
   in, re-bind the workspace TOC. Fast path reads the tagged bodies; if the tags
   were stripped (e.g. by Word), the repair path re-detects sections by their
   canonical heading text, syncs the recovered bodies into brief_sections, and
   re-renders a clean locked DOCX from the DB so the scaffold is restored.

These locks are honored by the OnlyOffice editor UI -- the in-app guarantee --
not cryptographically; the reattach pass is the recovery net for out-of-band edits.
"""
from __future__ import annotations

import re
import logging

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

logger = logging.getLogger(__name__)

# Canonical TRAP 38.1 catalog (key, title, word_excluded). Mirrors the catalog
# in modules/depositions/routes/brief_workspace_api.py -- kept local so this
# job module has no dependency on the route layer.
SECTION_CATALOG = [
    ("identity_of_parties", "Identity of Parties and Counsel", True),
    ("table_of_contents", "Table of Contents", True),
    ("index_of_authorities", "Index of Authorities", True),
    ("statement_of_the_case", "Statement of the Case", True),
    ("statement_on_oral_argument", "Statement Regarding Oral Argument", True),
    ("issues_presented", "Issues Presented", True),
    ("statement_of_facts", "Statement of Facts", False),
    ("summary_of_the_argument", "Summary of the Argument", False),
    ("argument", "Argument", False),
    ("prayer", "Prayer", False),
    ("certifications", "Certifications", True),
    ("appendix", "Appendix", True),
]

# Sections whose bodies are machine-generated on assembly -- locked entirely.
GENERATED = {"table_of_contents", "index_of_authorities", "certifications"}

SECTION_TITLES = {k: t for (k, t, _e) in SECTION_CATALOG}
_TITLE_TO_KEY = {}


def _norm(s: str) -> str:
    """Normalize a heading for matching: lowercase, alnum+space only, collapsed."""
    s = (s or "").lower()
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


for _k, _t, _e in SECTION_CATALOG:
    _TITLE_TO_KEY[_norm(_t)] = _k


# --- content-control construction ------------------------------------------

def _make_sdt(tag: str, alias: str, lock_val: str, sdt_id: int):
    """Build an empty block-level <w:sdt> (returns the sdt element)."""
    sdt = OxmlElement("w:sdt")
    pr = OxmlElement("w:sdtPr")
    if alias:
        a = OxmlElement("w:alias"); a.set(qn("w:val"), alias); pr.append(a)
    t = OxmlElement("w:tag"); t.set(qn("w:val"), tag); pr.append(t)
    i = OxmlElement("w:id"); i.set(qn("w:val"), str(sdt_id)); pr.append(i)
    lk = OxmlElement("w:lock"); lk.set(qn("w:val"), lock_val); pr.append(lk)
    sdt.append(pr)
    sdt.append(OxmlElement("w:sdtContent"))
    return sdt


def _wrap(paras, sdt):
    """Move a contiguous run of Paragraphs (already in the body, in order) into
    the sdt's content, leaving the sdt in their original position."""
    if not paras:
        return
    content = sdt.find(qn("w:sdtContent"))
    paras[0]._p.addprevious(sdt)
    for p in paras:
        content.append(p._p)


def wrap_section(doc, key, title, heading_para, body_paras, generated, sdt_id):
    """Wrap a freshly-added section (its heading paragraph + body paragraphs)
    into locked content controls. Heading: edit+delete locked. Body: delete
    locked (editable) unless generated, in which case edit+delete locked."""
    hsdt = _make_sdt(key + "__h", title, "sdtContentLocked", sdt_id)
    _wrap([heading_para], hsdt)

    lock = "sdtContentLocked" if generated else "sdtLocked"
    bsdt = _make_sdt(key, title, lock, sdt_id + 1)
    if not body_paras:
        body_paras = [doc.add_paragraph()]
    _wrap(body_paras, bsdt)


# --- reading an existing DOCX back ------------------------------------------

def _para_text(p_el) -> str:
    return "".join((t.text or "") for t in p_el.iter(qn("w:t")))


def parse_tagged_sections(docx_path) -> dict:
    """Fast path: pull {section_key: body_text} from the tagged body content
    controls (tag == section_key; heading controls tagged '<key>__h' skipped)."""
    from docx import Document
    doc = Document(docx_path)
    out = {}
    for sdt in doc.element.body.iter(qn("w:sdt")):
        pr = sdt.find(qn("w:sdtPr"))
        if pr is None:
            continue
        tag_el = pr.find(qn("w:tag"))
        if tag_el is None:
            continue
        tag = tag_el.get(qn("w:val")) or ""
        if not tag or tag.endswith("__h"):
            continue
        content = sdt.find(qn("w:sdtContent"))
        if content is None:
            continue
        texts = [_para_text(p) for p in content.findall(qn("w:p"))]
        out[tag] = "\n".join(texts).strip()
    return out


def detect_sections_by_heading(docx_path) -> dict:
    """Repair path: walk every paragraph in document order and segment the body
    by canonical heading text. Paragraphs before the first recognized heading
    (the caption) are ignored. Returns {section_key: body_text}."""
    from docx import Document
    doc = Document(docx_path)
    cur, acc = None, {}
    for p_el in doc.element.body.iter(qn("w:p")):
        txt = _para_text(p_el)
        key = _TITLE_TO_KEY.get(_norm(txt))
        if key:
            cur = key
            acc.setdefault(key, [])
        elif cur is not None and txt.strip():
            acc[cur].append(txt)
    return {k: "\n".join(v).strip() for k, v in acc.items()}


# --- DB sync + reattach -----------------------------------------------------

def sync_sections_to_db(tenant, cid, mapping: dict) -> int:
    """Write recovered section bodies + refreshed word counts into brief_sections.
    Only the keys present in `mapping` are touched."""
    from modules.depositions.jobs.appellate_assemble import _connect
    n = 0
    conn = _connect()
    try:
        cur = conn.cursor()
        for key, text in mapping.items():
            wc = len(re.findall(r"\S+", text or ""))
            cur.execute(
                "UPDATE brief_sections SET body=%s, word_count=%s, updated_at=now() "
                "WHERE appellate_case_id=CAST(%s AS uuid) AND section_key=%s "
                "  AND TRIM(tenant_id)=%s",
                (text, wc, str(cid), key, (tenant or "").strip()))
            n += cur.rowcount or 0
        conn.commit()
    finally:
        conn.close()
    return n


def reattach_brief_toc(tenant, cid, incoming_docx_path, out_dir=None) -> dict:
    """Reattach the workspace TOC to an edited DOCX brought back in.

    1. Recover section bodies (tagged fast path, else heading-detection repair).
    2. Sync the recovered bodies into brief_sections (DB = canonical store).
    3. Re-render a clean, fully-locked DOCX from the DB so the scaffold returns.
    """
    draftable = [k for (k, _t, _e) in SECTION_CATALOG if k not in GENERATED]

    tagged = parse_tagged_sections(incoming_docx_path)
    have = [k for k in draftable if k in tagged]
    repaired = False
    if len(have) >= max(1, (len(draftable) + 1) // 2):
        mapping = {k: tagged[k] for k in have}
    else:
        detected = detect_sections_by_heading(incoming_docx_path)
        mapping = {k: v for k, v in detected.items() if k in draftable}
        repaired = True

    matched = sorted(mapping.keys())
    unmatched = [k for k in draftable if k not in mapping]
    synced = sync_sections_to_db(tenant, cid, mapping)

    from modules.depositions.jobs.appellate_assemble import build_brief_docx
    built = build_brief_docx(tenant, cid, out_dir=out_dir)
    built.update({
        "reattached": True,
        "repaired": repaired,
        "matched": matched,
        "unmatched": unmatched,
        "synced_sections": synced,
    })
    logger.info("reattach_brief_toc: cid=%s repaired=%s matched=%s unmatched=%s",
                cid, repaired, matched, unmatched)
    return built
