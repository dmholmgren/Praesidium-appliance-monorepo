"""Report service — pre-built billing reports, Excel export."""
import json
from datetime import date
from decimal import Decimal
from typing import Optional
from io import BytesIO
from core.db.base import TenantSession
from modules.billing.models.report_definition import ReportDefinition

SYSTEM_REPORTS = [
    {"name": "Slip Listing - Detailed", "slug": "slip-listing-detailed", "category": "slip",
     "sql_template": "SELECT te.entry_date, te.description, te.hours, te.rate, te.amount, te.status, te.utbms_code, te.source, m.matter_number, m.matter_name, c.client_name, te.user_id FROM time_entries te JOIN matters m ON te.matter_id = m.id AND m.tenant_id = :tenant_id JOIN clients c ON m.client_id = c.id AND c.tenant_id = :tenant_id WHERE te.tenant_id = :tenant_id AND te.entry_date BETWEEN :date_from AND :date_to {attorney_filter} {matter_filter} {client_filter} ORDER BY te.entry_date",
     "column_definitions": json.dumps([{"key":"entry_date","label":"Date","type":"date"},{"key":"client_name","label":"Client","type":"text"},{"key":"matter_number","label":"Matter #","type":"text"},{"key":"description","label":"Description","type":"text"},{"key":"hours","label":"Hours","type":"decimal"},{"key":"rate","label":"Rate","type":"currency"},{"key":"amount","label":"Amount","type":"currency"},{"key":"status","label":"Status","type":"text"}]),
     "default_filters": json.dumps(["date_from","date_to","attorney_id","matter_id","client_id"])},
    {"name": "WIP Listing", "slug": "wip-listing", "category": "wip",
     "sql_template": "SELECT te.entry_date, te.description, te.hours, te.rate, te.amount, te.status, m.matter_number, m.matter_name, c.client_name FROM time_entries te JOIN matters m ON te.matter_id = m.id AND m.tenant_id = :tenant_id JOIN clients c ON m.client_id = c.id AND c.tenant_id = :tenant_id WHERE te.tenant_id = :tenant_id AND te.status IN ('draft','ai_suggested','reviewed','approved') AND te.invoice_id IS NULL {attorney_filter} {client_filter} ORDER BY te.entry_date DESC",
     "column_definitions": json.dumps([{"key":"entry_date","label":"Date","type":"date"},{"key":"client_name","label":"Client","type":"text"},{"key":"matter_number","label":"Matter #","type":"text"},{"key":"description","label":"Description","type":"text"},{"key":"hours","label":"Hours","type":"decimal"},{"key":"amount","label":"Amount","type":"currency"}]),
     "default_filters": json.dumps(["attorney_id","client_id"])},
    {"name": "WIP Aging", "slug": "wip-aging", "category": "wip",
     "sql_template": "SELECT c.client_name, m.matter_number, SUM(CASE WHEN DATEDIFF(CURDATE(), te.entry_date) <= 30 THEN te.amount ELSE 0 END) AS days_0_30, SUM(CASE WHEN DATEDIFF(CURDATE(), te.entry_date) BETWEEN 31 AND 60 THEN te.amount ELSE 0 END) AS days_31_60, SUM(CASE WHEN DATEDIFF(CURDATE(), te.entry_date) BETWEEN 61 AND 90 THEN te.amount ELSE 0 END) AS days_61_90, SUM(CASE WHEN DATEDIFF(CURDATE(), te.entry_date) > 90 THEN te.amount ELSE 0 END) AS days_over_90, SUM(te.amount) AS total_wip FROM time_entries te JOIN matters m ON te.matter_id = m.id AND m.tenant_id = :tenant_id JOIN clients c ON m.client_id = c.id AND c.tenant_id = :tenant_id WHERE te.tenant_id = :tenant_id AND te.invoice_id IS NULL AND te.status IN ('draft','ai_suggested','reviewed','approved') {client_filter} GROUP BY c.client_name, m.matter_number ORDER BY total_wip DESC",
     "column_definitions": json.dumps([{"key":"client_name","label":"Client","type":"text"},{"key":"matter_number","label":"Matter #","type":"text"},{"key":"days_0_30","label":"0-30","type":"currency"},{"key":"days_31_60","label":"31-60","type":"currency"},{"key":"days_61_90","label":"61-90","type":"currency"},{"key":"days_over_90","label":"90+","type":"currency"},{"key":"total_wip","label":"Total WIP","type":"currency"}]),
     "default_filters": json.dumps(["client_id"])},
    {"name": "Outstanding Invoices", "slug": "outstanding-invoices", "category": "invoice",
     "sql_template": "SELECT i.invoice_number, i.invoice_date, i.due_date, c.client_name, i.total_amount, i.balance_due, i.status, DATEDIFF(CURDATE(), i.due_date) AS days_overdue FROM invoices i JOIN clients c ON i.client_id = c.id AND c.tenant_id = :tenant_id WHERE i.tenant_id = :tenant_id AND i.balance_due > 0 AND i.status NOT IN ('void') {client_filter} ORDER BY i.due_date",
     "column_definitions": json.dumps([{"key":"invoice_number","label":"Invoice #","type":"text"},{"key":"invoice_date","label":"Date","type":"date"},{"key":"due_date","label":"Due","type":"date"},{"key":"client_name","label":"Client","type":"text"},{"key":"total_amount","label":"Total","type":"currency"},{"key":"balance_due","label":"Balance","type":"currency"},{"key":"days_overdue","label":"Days Overdue","type":"integer"}]),
     "default_filters": json.dumps(["client_id"])},
    {"name": "Receivables Aging", "slug": "receivables-aging", "category": "invoice",
     "sql_template": "SELECT c.client_name, SUM(CASE WHEN DATEDIFF(CURDATE(), i.due_date) <= 30 THEN i.balance_due ELSE 0 END) AS days_0_30, SUM(CASE WHEN DATEDIFF(CURDATE(), i.due_date) BETWEEN 31 AND 60 THEN i.balance_due ELSE 0 END) AS days_31_60, SUM(CASE WHEN DATEDIFF(CURDATE(), i.due_date) BETWEEN 61 AND 90 THEN i.balance_due ELSE 0 END) AS days_61_90, SUM(CASE WHEN DATEDIFF(CURDATE(), i.due_date) > 90 THEN i.balance_due ELSE 0 END) AS days_over_90, SUM(i.balance_due) AS total_ar FROM invoices i JOIN clients c ON i.client_id = c.id AND c.tenant_id = :tenant_id WHERE i.tenant_id = :tenant_id AND i.balance_due > 0 AND i.status NOT IN ('void') {client_filter} GROUP BY c.client_name ORDER BY total_ar DESC",
     "column_definitions": json.dumps([{"key":"client_name","label":"Client","type":"text"},{"key":"days_0_30","label":"0-30","type":"currency"},{"key":"days_31_60","label":"31-60","type":"currency"},{"key":"days_61_90","label":"61-90","type":"currency"},{"key":"days_over_90","label":"90+","type":"currency"},{"key":"total_ar","label":"Total AR","type":"currency"}]),
     "default_filters": json.dumps(["client_id"])},
    {"name": "Collections Report", "slug": "collections", "category": "collections",
     "sql_template": "SELECT p.payment_date, c.client_name, i.invoice_number, p.amount, p.method, p.reference_number FROM payments p JOIN invoices i ON p.invoice_id = i.id AND i.tenant_id = :tenant_id JOIN clients c ON i.client_id = c.id AND c.tenant_id = :tenant_id WHERE p.tenant_id = :tenant_id AND p.payment_date BETWEEN :date_from AND :date_to {client_filter} ORDER BY p.payment_date DESC",
     "column_definitions": json.dumps([{"key":"payment_date","label":"Date","type":"date"},{"key":"client_name","label":"Client","type":"text"},{"key":"invoice_number","label":"Invoice #","type":"text"},{"key":"amount","label":"Amount","type":"currency"},{"key":"method","label":"Method","type":"text"}]),
     "default_filters": json.dumps(["date_from","date_to","client_id"])},
    {"name": "Timekeeper Productivity", "slug": "timekeeper-productivity", "category": "distribution",
     "sql_template": "SELECT te.user_id, SUM(te.hours) AS hours_billed, SUM(te.amount) AS value_billed, COALESCE(SUM(te.amount)/NULLIF(SUM(te.hours),0),0) AS effective_rate FROM time_entries te WHERE te.tenant_id = :tenant_id AND te.entry_date BETWEEN :date_from AND :date_to {attorney_filter} GROUP BY te.user_id ORDER BY value_billed DESC",
     "column_definitions": json.dumps([{"key":"user_id","label":"Timekeeper","type":"text"},{"key":"hours_billed","label":"Hours","type":"decimal"},{"key":"value_billed","label":"Value","type":"currency"},{"key":"effective_rate","label":"Effective Rate","type":"currency"}]),
     "default_filters": json.dumps(["date_from","date_to","attorney_id"])},
    {"name": "Matter Profitability", "slug": "matter-profitability", "category": "distribution",
     "sql_template": "SELECT m.matter_number, m.matter_name, c.client_name, SUM(te.hours) AS total_hours, SUM(te.amount) AS total_billed FROM time_entries te JOIN matters m ON te.matter_id = m.id AND m.tenant_id = :tenant_id JOIN clients c ON m.client_id = c.id AND c.tenant_id = :tenant_id WHERE te.tenant_id = :tenant_id AND te.entry_date BETWEEN :date_from AND :date_to {client_filter} GROUP BY m.matter_number, m.matter_name, c.client_name ORDER BY total_billed DESC",
     "column_definitions": json.dumps([{"key":"matter_number","label":"Matter #","type":"text"},{"key":"matter_name","label":"Matter","type":"text"},{"key":"client_name","label":"Client","type":"text"},{"key":"total_hours","label":"Hours","type":"decimal"},{"key":"total_billed","label":"Billed","type":"currency"}]),
     "default_filters": json.dumps(["date_from","date_to","client_id"])},
    {"name": "Distribution Detail", "slug": "distribution-detail", "category": "distribution",
     "sql_template": "SELECT d.user_id, d.credit_type, d.matter_id, d.hours, d.gross_amount, d.overhead_deduction, d.net_amount, d.distribution_period, m.matter_number FROM distributions d JOIN matters m ON d.matter_id = m.id AND m.tenant_id = :tenant_id WHERE d.tenant_id = :tenant_id AND d.distribution_period = :period {attorney_filter} ORDER BY d.user_id",
     "column_definitions": json.dumps([{"key":"user_id","label":"Attorney","type":"text"},{"key":"credit_type","label":"Type","type":"text"},{"key":"matter_number","label":"Matter #","type":"text"},{"key":"hours","label":"Hours","type":"decimal"},{"key":"gross_amount","label":"Gross","type":"currency"},{"key":"overhead_deduction","label":"Overhead","type":"currency"},{"key":"net_amount","label":"Net","type":"currency"}]),
     "default_filters": json.dumps(["period","attorney_id"])},
    {"name": "Projected Distribution - WIP", "slug": "projected-distribution-wip", "category": "distribution",
     "sql_template": "SELECT te.user_id, SUM(te.hours) AS wip_hours, SUM(te.amount) AS wip_value, SUM(te.amount)*0.65 AS projected_net FROM time_entries te WHERE te.tenant_id = :tenant_id AND te.invoice_id IS NULL AND te.status IN ('draft','ai_suggested','reviewed','approved') {attorney_filter} GROUP BY te.user_id ORDER BY wip_value DESC",
     "column_definitions": json.dumps([{"key":"user_id","label":"Attorney","type":"text"},{"key":"wip_hours","label":"WIP Hours","type":"decimal"},{"key":"wip_value","label":"WIP Value","type":"currency"},{"key":"projected_net","label":"Projected Net","type":"currency"}]),
     "default_filters": json.dumps(["attorney_id"])},
]


class ReportService:
    def __init__(self, db: TenantSession):
        self.db = db

    async def seed_system_reports(self):
        existing = self.db.query(ReportDefinition).filter(ReportDefinition.is_system == True).count()
        if existing > 0:
            return
        for idx, rpt in enumerate(SYSTEM_REPORTS):
            rd = ReportDefinition(tenant_id=self.db.tenant_id, name=rpt["name"], slug=rpt["slug"],
                category=rpt["category"], sql_template=rpt["sql_template"],
                default_filters=rpt.get("default_filters"), column_definitions=rpt.get("column_definitions"),
                sort_order=idx, is_system=True)
            self.db.add(rd)
        self.db.commit()

    async def list_reports(self, category=None) -> list:
        query = self.db.query(ReportDefinition).filter(ReportDefinition.is_active == True)
        if category:
            query = query.filter(ReportDefinition.category == category)
        return query.order_by(ReportDefinition.sort_order).all()

    async def run_report(self, slug: str, filters: dict) -> dict:
        from sqlalchemy import text
        report_def = self.db.query(ReportDefinition).filter(ReportDefinition.slug == slug, ReportDefinition.is_active == True).first()
        if not report_def:
            raise ValueError(f"Report not found: {slug}")
        sql = report_def.sql_template
        params = {"tenant_id": self.db.tenant_id}
        replacements = {"attorney_filter": "", "matter_filter": "", "client_filter": "", "status_filter": ""}
        if filters.get("attorney_id"):
            replacements["attorney_filter"] = "AND te.user_id = :attorney_id"
            params["attorney_id"] = filters["attorney_id"]
        if filters.get("matter_id"):
            replacements["matter_filter"] = "AND te.matter_id = :matter_id"
            params["matter_id"] = filters["matter_id"]
        if filters.get("client_id"):
            replacements["client_filter"] = "AND c.id = :client_id"
            params["client_id"] = filters["client_id"]
        for k in ("date_from", "date_to", "period"):
            if filters.get(k):
                params[k] = filters[k]
        for placeholder, replacement in replacements.items():
            sql = sql.replace(f"{{{placeholder}}}", replacement)
        result = self.db.execute(text(sql), params)
        columns = list(result.keys())
        rows = [dict(zip(columns, row)) for row in result.fetchall()]
        for row in rows:
            for key, val in row.items():
                if isinstance(val, Decimal):
                    row[key] = float(val)
        col_defs = json.loads(report_def.column_definitions) if report_def.column_definitions else []
        return {"report_name": report_def.name, "columns": col_defs, "rows": rows, "row_count": len(rows)}

    async def export_to_excel(self, slug: str, filters: dict) -> bytes:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Border, Side
        data = await self.run_report(slug, filters)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = data["report_name"][:31]
        hdr_font = Font(bold=True)
        hdr_fill = PatternFill(start_color="D5E8F0", end_color="D5E8F0", fill_type="solid")
        thin = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
        for col_idx, col_def in enumerate(data["columns"], 1):
            cell = ws.cell(row=1, column=col_idx, value=col_def["label"])
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.border = thin
        for row_idx, row in enumerate(data["rows"], 2):
            for col_idx, col_def in enumerate(data["columns"], 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row.get(col_def["key"], ""))
                cell.border = thin
                if col_def.get("type") == "currency":
                    cell.number_format = '$#,##0.00'
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
