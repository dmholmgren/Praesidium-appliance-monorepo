"""Distribution engine — eat-what-you-kill, BigInteger IDs."""
from datetime import datetime, date
from decimal import Decimal
from typing import Optional
from io import BytesIO
from core.db.base import TenantSession
from core.audit import write_audit
from core.models.billing import Distribution
from core.models.billing import Payment
from core.models.billing import Invoice
from modules.billing.models.invoice_line_item import InvoiceLineItem
from core.models.matter import Matter
from core.models.billing import TimeEntry


class DistributionService:
    def __init__(self, db: TenantSession):
        self.db = db

    async def calculate_distribution(self, payment_id: int, overhead_rate: Decimal = Decimal("0.35")) -> list:
        payment = self.db.query(Payment).filter(Payment.id == payment_id).first()
        if not payment:
            raise ValueError("Payment not found")
        invoice = self.db.query(Invoice).filter(Invoice.id == payment.invoice_id).first()
        if not invoice:
            raise ValueError("Invoice not found")
        distributions = []
        period = date.today().strftime("%Y-%m")
        for inv_matter in invoice.matters:
            matter = self.db.query(Matter).filter(Matter.id == inv_matter.matter_id).first()
            if not matter or invoice.total_amount <= 0:
                continue
            matter_share = (inv_matter.subtotal / invoice.total_amount) * payment.amount
            line_items = self.db.query(InvoiceLineItem).filter(
                InvoiceLineItem.invoice_id == invoice.id, InvoiceLineItem.matter_id == matter.id,
                InvoiceLineItem.time_entry_id.isnot(None)).all()
            total_hours = sum(float(li.hours or 0) for li in line_items)
            working_by_user = {}
            for li in line_items:
                if li.time_entry_id:
                    entry = self.db.query(TimeEntry).filter(TimeEntry.id == li.time_entry_id).first()
                    if entry:
                        uid = entry.user_id
                        if uid not in working_by_user:
                            working_by_user[uid] = {"hours": Decimal("0"), "amount": Decimal("0")}
                        working_by_user[uid]["hours"] += entry.hours
                        working_by_user[uid]["amount"] += li.amount
            for uid, data in working_by_user.items():
                proportion = float(data["hours"]) / total_hours if total_hours > 0 else 0
                gross = matter_share * Decimal(str(proportion))
                overhead = gross * overhead_rate
                net = gross - overhead
                dist = Distribution(tenant_id=self.db.tenant_id, payment_id=payment_id,
                    invoice_id=invoice.id, matter_id=matter.id, user_id=uid, credit_type="working",
                    hours=data["hours"], rate=data["amount"] / data["hours"] if data["hours"] > 0 else Decimal("0"),
                    gross_amount=gross, overhead_deduction=overhead, net_amount=net, distribution_period=period)
                self.db.add(dist)
                distributions.append(dist)
            if matter.originating_attorney_id:
                orig_gross = matter_share * Decimal("0.10")
                orig_overhead = orig_gross * overhead_rate
                dist = Distribution(tenant_id=self.db.tenant_id, payment_id=payment_id,
                    invoice_id=invoice.id, matter_id=matter.id, user_id=matter.originating_attorney_id,
                    credit_type="origination", gross_amount=orig_gross, overhead_deduction=orig_overhead,
                    net_amount=orig_gross - orig_overhead, distribution_period=period)
                self.db.add(dist)
                distributions.append(dist)
        write_audit(self.db, "distribution.calculate", "payment", payment_id, new_values={"distributions_created": len(distributions, user_id=0)})
        self.db.commit()
        return distributions

    async def run_monthly_distribution(self, period_start: date, period_end: date, overhead_rate=Decimal("0.35")) -> dict:
        payments = self.db.query(Payment).filter(Payment.payment_date >= period_start, Payment.payment_date <= period_end).all()
        total = 0
        for payment in payments:
            existing = self.db.query(Distribution).filter(Distribution.payment_id == payment.id).first()
            if not existing:
                dists = await self.calculate_distribution(payment.id, overhead_rate)
                total += len(dists)
        return {"payments_processed": len(payments), "distributions_created": total}

    async def generate_distribution_report(self, period: str, user_id: Optional[int] = None) -> bytes:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Border, Side
        query = self.db.query(Distribution).filter(Distribution.distribution_period == period)
        if user_id:
            query = query.filter(Distribution.user_id == user_id)
        distributions = query.order_by(Distribution.user_id, Distribution.credit_type).all()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Distribution {period}"
        hdr_font = Font(bold=True, size=11)
        hdr_fill = PatternFill(start_color="D5E8F0", end_color="D5E8F0", fill_type="solid")
        thin = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
        headers = ["Attorney ID", "Credit Type", "Matter ID", "Hours", "Rate", "Gross", "Overhead", "Net"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.border = thin
        for row_idx, d in enumerate(distributions, 2):
            vals = [d.user_id, d.credit_type, d.matter_id, float(d.hours or 0), float(d.rate or 0), float(d.gross_amount), float(d.overhead_deduction), float(d.net_amount)]
            for col, v in enumerate(vals, 1):
                cell = ws.cell(row=row_idx, column=col, value=v)
                cell.border = thin
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
