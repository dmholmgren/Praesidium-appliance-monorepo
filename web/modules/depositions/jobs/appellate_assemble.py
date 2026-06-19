"""appellate_assemble.py -- Module B / Unit 10: assemble the brief -> TRAP-formatted
DOCX + PDF, then promote into the matter DMS.

Compiles the brief_sections (U6) in TRAP 38.1 order into a filed-format DOCX:
14-pt Times body (TRAP 9.4(e)), centered caption, a Word Table-of-Contents field,
an Index of Authorities compiled from the brief's own cites (U5 engine), and a
Certificate of Compliance carrying the live counted-word total (TRAP 9.4(i)). PDF
via the soffice lane (the depo-report converter). On --promote the DOCX+PDF are
filed into the matter's DMS folder (documents rows) via the drafting_outputs path.

CLI (inside praesidium-web, cwd /app):
  python -m modules.depositions.jobs.appellate_assemble --appeal UUID [--promote]
        [--subfolder "03-Briefs and Motions"] [--tenant T]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import logging
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile

logger = logging.getLogger(__name__)

GENERATED = {"table_of_contents", "index_of_authorities", "certifications"}
_TOA_GROUP = {"case": "Cases", "statute": "Statutes", "rule": "Rules",
              "constitution": "Constitutional", "regulation": "Other"}


def _connect():
    import psycopg2
    from modules.ediscovery.jobs.geometry_intake import _db_kwargs
    conn = psycopg2.connect(**_db_kwargs())
    conn.autocommit = False
    return conn


def _toc_field(doc):
    """Insert a Word TOC field (updates in Word; placeholder text in the meantime)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, txt in (("begin", None), ("instrText", 'TOC \\o "1-2" \\h \\z \\u'),
                      ("separate", None), ("text", "Right-click and Update Field in Word."),
                      ("end", None)):
        if kind == "instrText":
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        elif kind == "text":
            e = OxmlElement("w:t"); e.text = txt
        else:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
        run._r.append(e)
    return p


def _authorities_from_text(text):
    """Group + alphabetize the authorities cited in the assembled body (U5 engine)."""
    from modules.ediscovery.services.structural_pass import find_citations
    from modules.depositions.jobs.appellate_brief import _case_name_before, _paren_after
    agg = {}
    for c in find_citations(text):
        norm = c["norm"]
        if norm in agg:
            continue
        ct = c["ctype"]
        name = _case_name_before(text, c["start"]) if ct == "case" else None
        paren = _paren_after(text, c["end"]) if ct == "case" else None
        label = norm if not name else "%s, %s" % (name, norm)
        if paren:
            label += " " + paren
        agg[norm] = {"group": _TOA_GROUP.get(ct, "Other"), "label": label,
                     "sort": (name or norm).lower().lstrip("the ")}
    grouped = {}
    for rec in sorted(agg.values(), key=lambda r: (r["group"], r["sort"])):
        grouped.setdefault(rec["group"], []).append(rec["label"])
    return grouped


def build_brief_docx(tenant_id, appellate_case_id, out_dir=None) -> dict:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
    tenant = (tenant_id or "").strip()
    conn = _connect()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT matter_id::text, style, court_of_appeals, appellate_cause_number, "
            "       trial_court, trial_cause_number, appellant, appellee, jurisdiction "
            "FROM appellate_cases WHERE id=CAST(%s AS uuid) AND TRIM(tenant_id)=%s",
            (str(appellate_case_id), tenant))
        row = cur.fetchone()
        if not row:
            raise ValueError("appellate case not found")
        (mid, style, coa, appno, tcourt, tcause, appellant, appellee, jur) = row
        cur.execute(
            "SELECT section_key, title, body, word_excluded, included, word_count "
            "FROM brief_sections WHERE appellate_case_id=CAST(%s AS uuid) "
            "ORDER BY sort_order", (str(appellate_case_id),))
        sections = cur.fetchall()
    finally:
        conn.close()

    counted = sum(s[5] for s in sections if s[4] and not s[3])
    body_text = "\n\n".join((s[2] or "") for s in sections if s[4])
    authorities = _authorities_from_text(body_text)

    doc = Document()
    normal = doc.styles["Normal"].font
    normal.name = "Times New Roman"
    normal.size = Pt(14)                          # TRAP 9.4(e): 14-pt body

    def heading(txt):
        p = doc.add_paragraph()
        p.alignment = AL.CENTER
        r = p.add_run(txt.upper()); r.bold = True; r.font.size = Pt(14)
        return p

    def center(txt, bold=False):
        p = doc.add_paragraph(); p.alignment = AL.CENTER
        r = p.add_run(txt); r.bold = bold
        return p

    # --- caption ---
    center("IN THE " + (coa or "COURT OF APPEALS").upper())
    center("No. " + (appno or "______________"))
    center("")
    center((appellant or "Appellant") + ",", bold=True)
    center("Appellant,")
    center("v.")
    center((appellee or "Appellee") + ",", bold=True)
    center("Appellee.")
    center("")
    if tcourt or tcause:
        center("On Appeal from the " + (tcourt or "trial court") +
               (", Cause No. " + tcause if tcause else ""))
    center("")
    center("BRIEF OF APPELLANT", bold=True)
    doc.add_page_break()

    # --- sections (each wrapped in locked content controls so the TRAP 38.1
    #     scaffold cannot be deleted/reordered in OnlyOffice; bodies stay
    #     editable and are tagged with section_key for TOC reattachment) ---
    from modules.depositions.jobs import appellate_toc_bind as _tb
    _sid = 1000
    for (key, title, body, excl, included, wcount) in sections:
        if not included:
            continue
        hp = heading(title)
        paras = []
        if key == "table_of_contents":
            paras.append(_toc_field(doc))
        elif key == "index_of_authorities":
            if not authorities:
                paras.append(doc.add_paragraph("[No authorities cited yet.]"))
            for group in ["Cases", "Statutes", "Rules", "Constitutional", "Other"]:
                if group not in authorities:
                    continue
                gp = doc.add_paragraph(); gr = gp.add_run(group); gr.bold = True
                paras.append(gp)
                for label in authorities[group]:
                    paras.append(doc.add_paragraph(label))
        elif key == "certifications":
            body_txt = body or ""
            if body_txt.strip():
                for para in body_txt.split("\n"):
                    paras.append(doc.add_paragraph(para))
            paras.append(doc.add_paragraph(
                "Certificate of Compliance. This brief contains %d words in the "
                "sections counted under Tex. R. App. P. 9.4(i)(1), as computed by the "
                "word-processing software, and complies with the 15,000-word limit of "
                "Rule 9.4(i)(2)(B)." % counted))
            paras.append(doc.add_paragraph(
                "Certificate of Service. I certify that a true and correct copy of "
                "this brief was served on all counsel of record via the electronic "
                "filing manager on the date of filing."))
        else:
            if (body or "").strip():
                for para in body.split("\n"):
                    paras.append(doc.add_paragraph(para))
            else:
                paras.append(doc.add_paragraph("[To be drafted.]"))
        _tb.wrap_section(doc, key, title, hp, paras, key in GENERATED, _sid)
        _sid += 10

    out_dir = out_dir or tempfile.mkdtemp(prefix="brief_assemble_")
    safe = re.sub(r"[^A-Za-z0-9 .-]", "_", (style or "Brief"))[:80]
    fname = "Brief of Appellant - %s (Praesidium).docx" % safe
    docx_path = os.path.join(out_dir, fname)
    doc.save(docx_path)
    return {"docx_path": docx_path, "filename": fname, "matter_id": mid,
            "counted_words": counted, "authorities": sum(len(v) for v in authorities.values()),
            "sections": len([s for s in sections if s[4]])}


def to_pdf(docx_path) -> str:
    outdir = os.path.dirname(docx_path)
    prof = os.path.join(outdir, ".lo_profile")
    subprocess.run(["soffice", "--headless", "-env:UserInstallation=file://%s" % prof,
                    "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   check=True, capture_output=True, timeout=180)
    pdf_path = docx_path[:-5] + ".pdf"
    if not os.path.exists(pdf_path):
        raise RuntimeError("soffice produced no pdf")
    return pdf_path


def _file_into_dms(cur, tenant, matter_id, src_path, subfolder):
    """Copy a built file into the matter DMS folder + insert a documents row."""
    from pathlib import Path
    cur.execute("SELECT disk_root FROM matter_folders WHERE matter_id=CAST(%s AS uuid) "
                "AND TRIM(tenant_id)=%s AND disk_root LIKE '/mnt/praesidium%%' LIMIT 1",
                (str(matter_id), tenant))
    r = cur.fetchone()
    if not r:
        raise RuntimeError("no Praesidium folder for matter")
    dest_dir = Path(r[0]) / subfolder
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / Path(src_path).name
    stem, ext = dest.stem, dest.suffix
    i = 1
    while dest.exists():
        dest = dest_dir / ("%s (%d)%s" % (stem, i, ext)); i += 1
    shutil.copy2(src_path, str(dest))
    content = dest.read_bytes()
    checksum = hashlib.sha256(content).hexdigest()
    mime = mimetypes.guess_type(str(dest))[0] or "application/octet-stream"
    cur.execute(
        "INSERT INTO documents (id, tenant_id, matter_id, filename, original_filename, "
        "  title, mime_type, file_size, storage_path, document_type, status, checksum, "
        "  created_at, updated_at) VALUES (gen_random_uuid(), %s, CAST(%s AS uuid), %s, %s, "
        "  %s, %s, %s, %s, 'brief', 'active', %s, NOW(), NOW()) RETURNING id::text",
        (tenant, str(matter_id), dest.name, dest.name, dest.name, mime, len(content),
         str(dest), checksum))
    return cur.fetchone()[0], str(dest)


def assemble_brief(tenant_id, appellate_case_id, promote=False,
                   subfolder="03-Briefs and Motions") -> dict:
    tenant = (tenant_id or "").strip()
    built = build_brief_docx(tenant, appellate_case_id)
    pdf_path = to_pdf(built["docx_path"])
    out = {"filename": built["filename"], "docx_path": built["docx_path"],
           "pdf_path": pdf_path, "counted_words": built["counted_words"],
           "authorities": built["authorities"], "sections": built["sections"],
           "promoted": False}
    if promote:
        conn = _connect()
        try:
            cur = conn.cursor()
            ddoc, dpath = _file_into_dms(cur, tenant, built["matter_id"],
                                         built["docx_path"], subfolder)
            pdoc, ppath = _file_into_dms(cur, tenant, built["matter_id"], pdf_path, subfolder)
            # record the staged draft (drafting_outputs) as promoted
            cur.execute(
                "INSERT INTO drafting_outputs (tenant_id, matter_id, filename, storage_path, "
                "  document_type, file_size, draft_text, created_at, promoted_at, "
                "  promoted_path, promoted_doc_id) VALUES (%s, CAST(%s AS uuid), %s, %s, "
                "  'brief', %s, %s, NOW(), NOW(), %s, CAST(%s AS uuid))",
                (tenant, str(built["matter_id"]), built["filename"], dpath,
                 os.path.getsize(built["docx_path"]), "Assembled appellate brief draft",
                 dpath, ddoc))
            conn.commit()
            out.update({"promoted": True, "dms_docx_doc_id": ddoc, "dms_docx_path": dpath,
                        "dms_pdf_doc_id": pdoc, "dms_pdf_path": ppath})
        finally:
            conn.close()
    return out


def main():
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    ap = argparse.ArgumentParser(description="Assemble appellate brief (Module B / U10)")
    ap.add_argument("--tenant", default=os.environ.get("TENANT_ID", ""))
    ap.add_argument("--appeal", required=True)
    ap.add_argument("--promote", action="store_true")
    ap.add_argument("--subfolder", default="03-Briefs and Motions")
    args = ap.parse_args()
    out = assemble_brief(args.tenant, args.appeal, promote=args.promote, subfolder=args.subfolder)
    logger.info("DONE %s", json.dumps(out, default=str))


if __name__ == "__main__":
    main()
